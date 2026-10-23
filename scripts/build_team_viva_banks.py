"""Question + answer Word banks for Aarya, Akash, and Tanish."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

AARYA = [
    ("Q1", "What is the name of this project?",
     "An Interoperable Wearable Health Data Framework for Real-Time Predictive Health Monitoring and Early Disease Risk Detection. The demo site is called PulseVanguard AI."),
    ("Q2", "What problem does the paper start with?",
     "Care is still reactive. People are checked mainly after symptoms appear, or only at clinic visits, so chronic disease can worsen in the gaps."),
    ("Q3", "Name the three recurring technology gaps the survey isolates after reviewing twenty studies from 2016–2025.",
     "Fragmented wearable data across brands; no personalised health baseline; alerts that go to only one recipient at a time."),
    ("Q4", "In what way do proprietary wearable ecosystems (Apple, Samsung, Fitbit) create a clinical blind spot even when the sensors themselves are accurate?",
     "Each brand keeps data in its own silo. Accurate sensors still fail the clinic if the numbers never join one timeline the doctor can use."),
    ("Q5", "What does WHO say about non-communicable diseases?",
     "NCDs such as heart disease, diabetes, hypertension, and chronic lung disease cause about 74% of deaths worldwide, and many are preventable with earlier detection and monitoring."),
    ("Q6", "What does WHO’s 74% figure actually refer to?",
     "The share of global deaths from non-communicable diseases — not infections. We use it as motivation, not as a result our demo produced."),
    ("Q7", "Why are clinic visits not enough?",
     "A visit is a snapshot. Between appointments there is a blind spot. Wearables can fill it with a continuous stream — if the data is actually used."),
    ("Q8", "What are the 3 technical gaps found by the research paper?",
     "Same three: fragmented ecosystems, missing personal baseline / composite scoring, and single-recipient alerts. These are architecture gaps, not just ‘need a better neural net.’"),
    ("Q9", "What is HRI, simply?",
     "Health Risk Index: one 0–100 score of how concerning current vitals look. Bands: low 0–25, moderate 25–50, high 50–75, critical 75–100."),
    ("Q11", "Which databases were searched?",
     "IEEE Xplore, PubMed/MEDLINE, Springer, MDPI, plus IJARCCE, Elsevier health-informatics journals, and NEJM Catalyst."),
    ("Q12", "Why is a single abnormal vital (for example tachycardia) an insufficient basis for an early-warning claim?",
     "Heart rate jumps with stairs, stress, or coffee. Without SpO₂, BP, temperature, activity, and trend you get false alarms. Multi-sensor fusion is stronger; that is why HRI combines several vitals."),
    ("Q13", "Which five thematic areas structured the literature search, and why is that taxonomy better than grouping papers only by year?",
     "Wearable IoT monitoring; ML risk prediction; interoperability; smart alerting; virtual assistants. Themes map to the gaps we close. Year-only grouping would mix unrelated problems and hide that no one paper did all five."),
    ("Q14", "What’s a three-channel alert?",
     "The same event explained to patient, caregiver, and clinician together. In the demo those three messages show on the Alerts panel (not actually emailed/WhatsApped)."),
    ("Q15", "Is this a medical device?",
     "No. Academic educational prototype on synthetic data. Not a registered device, not a diagnosis, not a doctor."),
    ("Q16", "Say the ethical line before you use the words ‘early disease detection.’",
     "We are not detecting disease in a real body. HRI is a simulated score on synthetic vitals. The title is the research ambition, not a clinical proof from this demo."),
    ("Q17", "If they ask ‘can this replace doctors?’",
     "No. Teaching monitor and explainer. Decisions stay with clinicians. Our contribution is architecture, not autonomous medicine."),
    ("Q18", "Outline the seven-stage pipeline from wearable device to intervention.",
     "1) Devices collect. 2) APIs/connectors pull data. 3) Normalise to one schema. 4) AI/ML scores risk. 5) Composite HRI. 6) Three-channel alerts. 7) Dashboard and assistant (lifestyle engine is in the paper, not a full diet app in the live build)."),
    ("Q19", "Why is Random Forest described as more robust than SVM once real-world sensor noise is present?",
     "Zhou et al. saw RF and SVM both above ~85% in lab-style tasks, but RF held up better once realistic wearable noise was added. Noisy tables suit trees better than a clean-lab SVM."),
    ("Q20", "Why LSTM?",
     "LSTM remembers sequences over time — hourly history, not one reading. Miotto argued that for longitudinal records; Rigla used LSTM on glucose with ~30 min lead. RF looks at a moment; LSTM looks at a story. In our laptop demo LSTM is optional if TensorFlow is missing."),
    ("Q21", "Name one limitation of the current software.",
     "Pick any honest one: vendor OAuth is stubbed; chat can be slow on CPU; LSTM may be absent; data is synthetic; alerts are simulated on screen."),
    ("Q22", "How would you explain the project in 20 seconds?",
     "Watches already measure, but data sits in silos and care is late. We unify vitals, score HRI, warn patient, family, and clinician, and let an assistant explain those numbers. Demo only."),
    ("Q23", "Name three wearables the paper talks about.",
     "Apple Watch, Samsung Galaxy Watch, Fitbit."),
    ("Q24", "Kajornkasirat got high correlation with medical sensors. Why isn’t that the same as clinical decision support?",
     "High R means the gadget tracked a reference sensor. Decision support means you change care inside a workflow. They validated sensing, not a care pathway. We have not clinically validated HRI either."),
    ("Q25", "Table II contrasts custom IoT hardware with consumer APIs. What deployment and cost argument follows from choosing APIs?",
     "Custom kits mean you manufacture hardware. APIs mean people use watches they already own — low hardware cost, easier scale. Honesty: we did not finish real OAuth; we score a dataset as if that layer existed."),
    ("Q26", "Where in the pipeline would FHIR/HL7 mapping sit if this moved beyond an academic prototype?",
     "After normalisation: map our JSON (user, time, metric, value, unit, source) into FHIR Observation or HL7 so an EHR could ingest it. Next step, not in the repo."),
    ("Q27", "What limitation of Kajornkasirat [1] and Ekanath [2] alert designs does Gap 2 in Table III mark, and how does your demo close it on screen?",
     "Gap 2 is single-recipient alerts — staff or caregivers, not all three. Open a high/critical synthetic patient; the Alerts panel lists patient, caregiver, and clinician together."),
    ("Q28", "What ethical boundary must you state before using the phrase ‘early disease detection’ in a viva?",
     "Do not claim we diagnosed real patients. Synthetic HRI only. The phrase is the field’s goal, not this prototype’s clinical result."),
    ("Q30", "What does PulseVanguard AI actually implement today versus what Table II lists as a full lifestyle-recommendation engine?",
     "Today: unified vitals, HRI, forecast, three alert texts, dashboard, Ollama chat. Lifestyle engine is listed in the paper/table; the live site does not run a full diet coach. Say: proposed, not fully implemented."),
    ("Q31", "How do you answer: ‘Can this system replace a physician?’ without undermining the contribution?",
     "It must not replace a physician. The contribution is combining HRI, three-channel alerts, and a live-data assistant in one academic stack. That is still valuable without claiming to practise medicine."),
    ("Q32", "If asked for originality, which combination of HRI + three-channel alerts + live assistant is hard to find in any single reviewed paper?",
     "Papers do IoT, or RF, or a bot, or hospital EWS — usually one slice. Hard to find in one paper: multi-source vitals, composite HRI, simultaneous three-role alerts, and an assistant that sees live history and forecast. That combo is our originality — not ‘we invented AI healthcare.’"),
]

AKASH = [
    ("Q1", "In what sense does the paper argue that healthcare still ‘operates reactively,’ and what temporal blind spot does that create between visits?",
     "Reactive means we mostly act after symptoms or at scheduled visits. Between visits nobody is holding a continuous strip, so hypertension or similar can drift for weeks unseen. That interval is the temporal blind spot."),
    ("Q2", "Why do the authors claim most NCD deaths are preventable, and how does that public-health claim license an engineering project?",
     "WHO links a large share of NCD deaths to late detection and weak ongoing monitoring. If earlier continuous signals can help, then software that unifies wearables, scores risk, and alerts people is a legitimate engineering response — not a claim that our demo already saved lives."),
    ("Q3", "Gupta and Sharma [3] achieved remote monitoring without an AI prediction layer. Why is that design now considered an incomplete solution?",
     "Streaming HR and temperature is necessary but not sufficient. You can watch a live graph and still have no idea if tomorrow looks worse. The missing middle is prediction and a shared score — which is why we added HRI, not only charts."),
    ("Q4", "Why mention Apple, Samsung, and Fitbit in the introduction?",
     "To show sensors are already on wrists. The remaining hole is a shared intelligent software layer, not inventing another gadget from scratch."),
    ("Q5", "What is a personal baseline?",
     "What is normal for this person. A runner’s resting HR is not an 80-year-old’s. Fixed cut-offs false-alarm some people and miss others. The paper wants per-user HRI; the demo still uses population bands plus that person’s history."),
    ("Q6", "Which Python app serves the website?",
     "FastAPI in backend/main.py, run with uvicorn on 127.0.0.1 port 8000."),
    ("Q7", "What is Streamlit used for?",
     "An optional second dashboard (frontend/app.py) for us as developers. The viva face is the PulseVanguard website. Same risk ideas, different UI."),
    ("Q8", "What if LSTM is missing?",
     "The demo does not crash. Random Forest (if weights exist) and/or the rule-based heuristic still produce HRI. LSTM is optional when TensorFlow/weights are absent."),
    ("Q9", "What is wrong with treating ‘user satisfaction 95%’ in a small prototype [2] as evidence of clinical effectiveness?",
     "Satisfaction in a tiny simulated prototype is not accuracy, not outcome, not a trial. Ekanath’s 95% is a usability headline, not proof we should treat patients. We must not quote it as clinical effectiveness."),
    ("Q10", "Why is single-lead consumer arrhythmia detection [10][11][13] still not a substitute for a multi-parameter risk index?",
     "Those papers are strong on ECG/AF or 14-class arrhythmia. Our problem is broader deterioration — HR, SpO₂, BP, temperature, respiration together. Catching AF is not the same as a composite HRI for chronic monitoring."),
    ("Q11", "State the problem as a systems-engineering mismatch: abundant sensors, scarce integrated intelligence.",
     "We have lots of watches and few systems that join the data, score risk, alert the right people, and explain it. Sensing is abundant; the missing product is the integration layer."),
    ("Q12", "Why is a single composite HRI preferable to firing independent alarms on each vital sign?",
     "Five separate beeps is alarm soup. One 0–100 score plus a band drives one policy: who is told and how strongly. Clinicians think ‘how worried are we about this person,’ not five uncoordinated toasts."),
    ("Q13", "Explain unit conversion and schema mapping as the concrete mechanism of interoperability in the demo.",
     "Each source may send kcal vs kJ, bpm vs interval, different field names. The normalizer maps everything into one record: user, timestamp, metric type, value, unit, source. That is interoperability in code, even when OAuth is still a stub."),
    ("Q14", "Why keep a heuristic HRI even after loading Random Forest weights?",
     "RF can under-score obviously bad vitals. Rules based on vital ranges still push HRI up. If the heuristic is already critical (≥75), a floor keeps the demo from looking ‘low risk’ during a crash scenario."),
    ("Q15", "Walk through FastAPI.",
     "FastAPI is the web backend. It serves index.html and demo.html, and JSON routes such as /patients, /patients/{id}/dashboard, /chat, /predict. The browser talks to 127.0.0.1:8000. You start it with uvicorn backend.main:app --reload --host 127.0.0.1 --port 8000."),
    ("Q16", "How would you justify 24-hour hourly history per synthetic patient as enough for a demo trend yet not enough for LSTM pretraining?",
     "Twenty-four points are enough to draw a slope, a small forecast, and a trend sentence for viva. LSTM wants long, dense sequences and large datasets (Miotto’s cold-start warning). We do not pretend 24 fake hours trained a world-class LSTM."),
    ("Q17", "What does the risk engine mix?",
     "RF probability if the pickle exists, LSTM if TensorFlow and weights exist, and a heuristic from vital ranges. Those are blended, then mapped to a band. Extreme heuristic can floor the score."),
    ("Q18", "What does the /chat API need?",
     "The user’s message and a patient_id so the server can attach vitals, age, gender, HRI, history text, trend, and forecast into the prompt."),
    ("Q19", "Why is the output of /chat AI slow?",
     "Ollama runs a local LLM on CPU. The first reply often loads the model into RAM (tens of seconds). Later replies are faster if keep_alive holds it. Do not click Send repeatedly."),
    ("Q20", "How does the dashboard’s trend sentence (‘rising heart rate; declining SpO₂’) communicate process to a clinician-examiner?",
     "It shows we are not only showing the latest row — we estimated slope over recent hours. That is the ‘process’ of longitudinal watching, even if the maths is a simple slope."),
    ("Q21", "What latency budget does a local CPU LLM chat have, and how do you set expectations?",
     "Alerts in papers like Ekanath are 3–5 seconds. Our chat is not that: first token can take up to a minute on CPU. We set timeout ~120s, cap reply length, and say ‘first answer is slow, then better.’ Alerts themselves are instant in-process messages."),
    ("Q22", "Name the four critical test identities used to stress alerts (severe hypoxia/tachycardia versus bradycardia/hypotension) without claiming they are real patients.",
     "Synthetic IDs P002, P010, P050, and P100. We pushed last hours to extreme simulated vitals (high HR/low SpO₂ versus low HR/low BP). They are not real people and not real diagnoses."),
    ("Q23", "If asked ‘what would you publish next?’, which gap (prospective wearable trial, FHIR, true personal baseline) is the honest sequel?",
     "Any of those is honest. Best: a small prospective wearable study with consent. FHIR export is the engineering sequel. True per-user baseline is what the paper promised more fully than the demo. Not ‘another chatbot skin.’"),
    ("Q24", "Why ‘zero hardware cost’ in the paper?",
     "The design uses consumer watches people already own instead of shipping a custom Arduino kit to every user. Cost of our college laptop demo is separate — we mean no extra clinical hardware in the proposed architecture."),
    ("Q25", "What is SpO₂?",
     "Blood oxygen saturation as a percentage. In the heuristic, normal is about 95–100%; values toward 90% or below raise HRI."),
    ("Q26", "Name one chatbot review result.",
     "Laranjo et al.: conversational agents can improve engagement (for example medication habits), but those reviewed bots were not hooked to live wearable streams."),
    ("Q27", "What language is the backend?",
     "Python."),
    ("Q28", "Give a 15-second process answer.",
     "Collect vitals, put them in one format, compute HRI with ML and rules, show three alerts, draw history and a short forecast, explain with a grounded chatbot."),
    ("Q29", "How do you keep demographics authoritative in chat (CSV age/gender) when a small local model prefers ‘not specified’ templates?",
     "We load age and gender from patient_demographics.csv, put them in the system and user prompt, and patch replies that still say ‘not specified.’ The file wins over the model’s bad habit."),
    ("Q30", "How do you present originality: first combination in this academic framing, not ‘first AI in healthcare’?",
     "Say: we are not first at AI. We are first in this student survey framing to put composite HRI, three-channel alerts, and a live-data assistant in one stack. Examiners respect that humility."),
    ("Q31", "What is the academic contribution in simple words?",
     "We reviewed 20 papers, named the gaps, and built a teaching system that shows how those gaps could be closed together — unify, score, alert three roles, explain. Not a hospital product."),
]

TANISH = [
    ("Q1", "How does PulseVanguard convert a reactive care model into a continuous preventive loop on the demo screen?",
     "You do not wait for a complaint. You pick a person, see current vitals and HRI, a six-hour sketch, and alerts if the score is high, then ask the bot. That is watch-ahead, still simulated."),
    ("Q2", "How does the implemented software explicitly address fragmentation, weak personalization, and single-recipient alerts?",
     "Fragmentation: one schema and one dashboard for all patients. Weak personalization: that person’s history, demographics, and HRI — not a generic page. Single-recipient: Alerts lists patient, caregiver, and clinician together."),
    ("Q3", "How does the implemented software explicitly address fragmentation, weak personalization, and single-recipient alerts?",
     "Same as Q2 — it is a duplicate on the list. Answer once: unified vitals, per-patient context, three-channel alerts on screen."),
    ("Q4", "How does alert-to-one-stakeholder design create asymmetric information among patient, family, and clinician?",
     "If only the patient sees it, they may ignore it. If only family gets WhatsApp, the doctor has no trail. If only the ICU sees it, the family is blind. Three channels are how we stop that split."),
    ("Q5", "If India-scale affordability is raised, how do consumer APIs and ‘zero custom hardware’ in Table III answer cost without pretending the demo is a national deployment?",
     "The paper’s cost story is: use watches people already bought, don’t manufacture kits. Our demo is a laptop at college, not a national rollout. Separate the architecture argument from ‘we deployed India-wide.’"),
    ("Q6", "How do you prevent the viva from sliding into ‘we built a chatbot’ instead of ‘we closed literature gaps’?",
     "Lead with the three gaps and HRI. Show alerts. Chat is last, and only to explain those numbers. If they only remember the bot, you under-sold the survey."),
    ("Q7", "Using the live demo, narrate one patient’s path from search/enrollment to vitals, HRI, forecast, alerts, and a grounded question to the assistant.",
     "Search Priya or P002 — or add a name/age/gender. Card fills. Read HRI and band. Scroll past readings. Forecast table. Alerts if high/critical. Ask ‘what is this patient’s HRI and age?’ Check it matches the card. Say the disclaimer."),
    ("Q8", "How is HRI computed when both Random Forest output and a rule-based heuristic are available?",
     "We mix RF’s score with a heuristic from vital ranges (about 45% RF / 55% rules when LSTM is off). If the heuristic is already ≥75, we take the max so a weak forest cannot hide a crash."),
    ("Q9", "How do you keep demographics authoritative in chat (CSV age/gender) when a small local model prefers ‘not specified’ templates?",
     "Inject age and gender from the demographics file into the prompt, and replace leftover ‘not specified’ text in the reply. The CSV is the source of truth."),
    ("Q10", "How do you tell the solution in one sentence?",
     "We unify vitals, compute a Health Risk Index, alert three roles, and let an assistant explain that patient’s numbers."),
    ("Q11", "What does ‘fragmented data’ look like to a user?",
     "Steps in one app, heart rate in another, clinic seeing none of it in one place."),
    ("Q12", "What screen proves prediction?",
     "The Forecast table: +hours, HR, SpO₂, predicted HRI, band."),
    ("Q13", "Is the assistant always correct?",
     "No. Small local models ramble. We cap length and attach facts. Always check the dashboard numbers."),
    ("Q14", "What is a good viva question to ask the bot?",
     "‘Summarise this patient’s trend and risk band’ or ‘Why might alerts have fired?’"),
    ("Q15", "What is a bad thing for the bot to say?",
     "A real diagnosis, a drug dose, or ‘go to hospital now’ as if it were a licensed clinician. We treat it as a demo explainer."),
    ("Q16", "What is patient_memory.py for?",
     "Load that person’s history, write a trend sentence, and build the short forecast."),
    ("Q17", "What is alert_system.py for?",
     "Turn an HRI band into patient / caregiver / clinician messages."),
    ("Q18", "What is ollama_agent.py for?",
     "Build the system prompt with vitals and history, call local Ollama, fall back if it is offline."),
    ("Q19", "What command trains models?",
     "From the project folder: python backend/models/train_model.py"),
    ("Q20", "What is CORS in one line?",
     "A browser rule we relax in the demo so the local page is allowed to call the local API."),
    ("Q21", "How does FastAPI serving website/demo.html plus JSON APIs illustrate the ‘dashboard and assistant’ stage of Table II?",
     "Table II’s last stage is visualisation plus a live-data assistant. FastAPI delivers the HTML dashboard and /chat JSON from the same app — one stage, two surfaces."),
    ("Q22", "Why mix RF probability with heuristic rather than taking argmax of either?",
     "Argmax would ignore half the evidence. Mixing uses ML when it is sensible and rules when physiology is obviously extreme. The floor on high heuristic is extra safety for the demo."),
    ("Q23", "What logging/privacy rule applies to API keys in user_store, and why mention it if connectors are stubbed?",
     "Keys are stored locally, encrypted at rest, not printed in logs, DB gitignored. Even stubs teach that you do not dump secrets into GitHub — good hygiene before real OAuth."),
    ("Q24", "Walk through what a critical HRI produces in the demo alert pane across the three channels.",
     "Band critical (≥75): patient sees an urgent rest-and-seek-help style message (still simulated). Caregiver gets an immediate check-in style note. Clinician gets a simulated escalation/triage line. All three appear together."),
    ("Q25", "How would you answer: ‘Is this a medical device?’ in one precise sentence?",
     "No — it is an educational software prototype on synthetic data and is not a registered medical device."),
    ("Q26", "What would a clinician still need to see that the current HRI does not encode (meds, labs, comorbidities)?",
     "Medications, lab results, diagnoses, symptoms, ECG. HRI is vital-based teaching score, not a full chart."),
    ("Q27", "How do you present originality: first combination in this academic framing, not ‘first AI in healthcare’?",
     "Humble line: first time we combined HRI + three alerts + live assistant in this survey-backed student framework — not the first AI in medicine."),
    ("Q28", "Give the solution as a closed loop: measure → normalize → score → notify three roles → explain with live context.",
     "Measure vitals, tidy them, score HRI, notify patient/caregiver/clinician, explain with history and forecast in the assistant. Then you can measure again."),
    ("Q29", "Why must lifestyle/diet recommendations remain modest or absent in the live build even though the paper lists a lifestyle engine?",
     "We did not validate diet advice. Over-claiming would be unsafe. Paper lists it as a proposed stage; the demo focuses on score, alerts, and Q&A."),
    ("Q30", "Mandl wanted what?",
     "A more open, patient-accessible health information economy through APIs — interoperability as the big barrier."),
    ("Q31", "Is this a medical device? One precise sentence.",
     "No; educational prototype, synthetic data, not a medical device. (Same as Q25.)"),
    ("Q32", "Two vitals that often lift HRI?",
     "Very high or very low heart rate, and SpO₂ falling toward or below about 90% in our rule table. Extreme BP or fever also contribute."),
    ("Q33", "Why did we use Python?",
     "Fast libraries: pandas for data, scikit-learn for Random Forest, FastAPI for the website, easy to demo in college."),
    ("Q34", "Easy difference vs Kajornkasirat?",
     "They built custom IoT hardware (MAX30102 etc.). We score a software dataset as if consumer wearables/APIs fed it, plus HRI dashboard and three alerts."),
    ("Q35", "How do you stop the viva becoming ‘we made a chatbot’?",
     "Same as Q6: start with gaps and HRI, bot last and grounded."),
    ("Q36", "What if the server’s down?",
     "Start uvicorn again. The 500 patients did not vanish from the CSV. Use http://127.0.0.1:8000/demo, not a raw HTML file."),
    ("Q37", "What did Kawamoto’s alert research teach?",
     "Clinical decision support helped in most trials when alerts fired automatically inside the workflow — not when staff had to remember to open a tool. That is why our alerts are automatic from HRI, not a hidden menu."),
    ("Q38", "What is multi-sensor fusion in plain words?",
     "Use several measurements together so one noisy sensor does not decide everything."),
    ("Q39", "What temperature is concerning in the heuristic table?",
     "Normal is about 36.1–37.2 °C. Toward 39 °C or below 35 °C is treated as more serious in the demo rules (critical_high 39.0, critical_low 35.0)."),
]


def _font(run, name="Calibri", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.95)
    sec.right_margin = Inches(0.95)


def _cover(doc: Document, title: str, person: str, kind: str) -> None:
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    _font(r, name="Georgia", size=22, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(f"{person}\n{kind}")
    _font(r, name="Georgia", size=14, italic=True, color=RGBColor(0x3D, 0x4B, 0xC8))
    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run(
        "Interoperable Wearable Health Data Framework  ·  PulseVanguard AI\n"
        "JSPM’s RSCOE, Pune  ·  Guide: Dr. Nilima Mohite  ·  AY 2025–26"
    )
    _font(r, size=10, color=RGBColor(0x4A, 0x50, 0x78))


def write_questions(path: Path, person: str, items: list) -> None:
    doc = Document()
    _setup(doc)
    _cover(doc, "Question Bank", person, "Questions only")
    n = doc.add_paragraph()
    r = n.add_run("Use this sheet in the hall. Practise answers from the matching Answer Bank.")
    _font(r, size=10, italic=True, color=RGBColor(0x4A, 0x50, 0x78))
    for num, q, _a in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(f"{num}. ")
        _font(r, size=12, bold=True, color=RGBColor(0x3D, 0x4B, 0xC8))
        r2 = p.add_run(q)
        _font(r2, size=12)
    doc.save(path)


def write_answers(path: Path, person: str, items: list) -> None:
    doc = Document()
    _setup(doc)
    _cover(doc, "Answer Bank", person, "Spoken-style answers")
    n = doc.add_paragraph()
    r = n.add_run("Always add: educational prototype, synthetic data, not a medical device.")
    _font(r, size=10, italic=True, color=RGBColor(0x4A, 0x50, 0x78))
    for num, q, a in items:
        pq = doc.add_paragraph()
        pq.paragraph_format.space_before = Pt(10)
        pq.paragraph_format.space_after = Pt(2)
        r = pq.add_run(f"{num}. ")
        _font(r, size=12, bold=True, color=RGBColor(0x3D, 0x4B, 0xC8))
        r2 = pq.add_run(q)
        _font(r2, size=12, bold=True)
        pa = doc.add_paragraph()
        pa.paragraph_format.space_after = Pt(4)
        pa.paragraph_format.left_indent = Inches(0.12)
        r = pa.add_run("A. ")
        _font(r, size=12, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
        r2 = pa.add_run(a)
        _font(r2, size=12)
    doc.save(path)


def write_combined() -> Path:
    doc = Document()
    _setup(doc)
    _cover(doc, "Team Answer Bank", "Aarya Thakare  ·  Akash Deshmukh  ·  Tanish Khedekar", "All viva questions with answers")
    n = doc.add_paragraph()
    r = n.add_run(
        "Tanish Q2 and Q3 are the same question. Q25 and Q31 are the medical-device line twice. "
        "Aarya Q3/Q8, Q16/Q28, Q17/Q31 overlap — say it twice in slightly different words."
    )
    _font(r, size=10, italic=True, color=RGBColor(0x4A, 0x50, 0x78))
    for person, items in (
        ("Aarya Thakare", AARYA),
        ("Akash Deshmukh", AKASH),
        ("Tanish Khedekar", TANISH),
    ):
        doc.add_page_break()
        h = doc.add_paragraph()
        r = h.add_run(person)
        _font(r, name="Georgia", size=16, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
        for num, q, a in items:
            pq = doc.add_paragraph()
            pq.paragraph_format.space_before = Pt(8)
            r = pq.add_run(f"{num}. ")
            _font(r, size=11, bold=True, color=RGBColor(0x3D, 0x4B, 0xC8))
            r2 = pq.add_run(q)
            _font(r2, size=11, bold=True)
            pa = doc.add_paragraph()
            pa.paragraph_format.left_indent = Inches(0.12)
            r = pa.add_run("A. ")
            _font(r, size=11, bold=True)
            r2 = pa.add_run(a)
            _font(r2, size=11)
    path = Path(r"d:\Team_Viva_Answer_Bank.docx")
    doc.save(path)
    return path


def main() -> None:
    pairs = [
        (r"d:\Aarya_Thakare_Question_Bank.docx", r"d:\Aarya_Thakare_Answer_Bank.docx", "Aarya Thakare", AARYA),
        (r"d:\Akash_Deshmukh_Question_Bank.docx", r"d:\Akash_Deshmukh_Answer_Bank.docx", "Akash Deshmukh", AKASH),
        (r"d:\Tanish_Khedekar_Question_Bank.docx", r"d:\Tanish_Khedekar_Answer_Bank.docx", "Tanish Khedekar", TANISH),
    ]
    for qp, ap, person, items in pairs:
        write_questions(Path(qp), person, items)
        write_answers(Path(ap), person, items)
        print("wrote", qp)
        print("wrote", ap)
    c = write_combined()
    print("wrote", c)


if __name__ == "__main__":
    main()
