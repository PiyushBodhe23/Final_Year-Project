"""Generate the 150-question viva bank as a Word document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SET_A = [
    # Problem (1–18) — includes paraphrased cores 1,2,3,8
    "Why is episodic, visit-based assessment structurally ill-suited to the management of chronic non-communicable disease?",
    "The introduction cites a WHO figure of roughly 74% of global deaths from NCDs. How does that statistic justify continuous wearable monitoring rather than more frequent clinic visits?",
    "Name the three recurring technology gaps the survey isolates after reviewing twenty studies from 2016–2025.",
    "In what way do proprietary wearable ecosystems (Apple, Samsung, Fitbit) create a clinical blind spot even when the sensors themselves are accurate?",
    "How do Mandl and Kohane [18] frame interoperability as a structural—not merely technical—barrier, and how does that claim motivate your API/normalization layer?",
    "Dias and Cunha [14] argue that hardware matured faster than software aggregation. What ‘multi-billion-dollar gap’ are they pointing to, and where does your framework sit relative to it?",
    "Why is a single abnormal vital (for example tachycardia) an insufficient basis for an early-warning claim?",
    "Distinguish reactive care, remote monitoring without prediction, and predictive preventive care as three different system ambitions.",
    "Kajornkasirat et al. [1] report high Pearson correlations against medical-grade devices. Why is correlation with a reference sensor still not equivalent to clinical decision support?",
    "Bassam et al. [9] showed earlier deterioration detection in COVID-19 home isolation. What must you not over-generalize from an infectious-disease cohort to chronic NCD surveillance?",
    "Escobar et al. [16] reduced ICU transfers using hospital EHR vitals. Why is that result encouraging yet incomplete as evidence for a consumer-wearable framework?",
    "If examiners ask ‘is fragmentation really unsolved in 2025?’, how do you answer using Deng et al. [6] on battery life and data standardization?",
    "What monitoring blind spot opens between scheduled outpatient visits, and how long can disease progress inside that interval?",
    "Why do the authors treat the largest unmet need as architectural rather than purely algorithmic?",
    "How would you state the problem in one sentence for a non-engineering examiner without using jargon?",
    "Which five thematic areas structured the literature search, and why is that taxonomy better than grouping papers only by year?",
    "What inclusion and exclusion rules governed the twenty-paper corpus, and what bias does a twenty-paper sample still carry?",
    "Why is ‘alert designed for one recipient’ a patient-safety issue and not only a notification-design issue?",
    # Process (19–34) — cores 4,5,10 paraphrased
    "Outline the seven-stage pipeline from wearable device to intervention as proposed in Section V.",
    "Define the Composite Health Risk Index (HRI) and state the four numerical bands implemented in software (low, moderate, high, critical).",
    "Why does Zhou et al. [5] plus Miotto [12] and Rigla et al. [15] jointly justify pairing a tree ensemble with a temporal model rather than using SVM alone?",
    "In the prototype, which vital fields enter the HRI feature vector, and why are steps included alongside physiology?",
    "Explain how the normalizer’s unified schema (timestamp, metric, unit, source, user) attacks vendor lock-in without requiring custom hardware.",
    "What is the role of a per-user physiological baseline, and how does a fixed population threshold fail elderly or athletic users?",
    "Walk through data flow in the repository: CSV/Excel history → FastAPI dashboard → risk engine → alert bundle → chat context.",
    "Why is Random Forest described as more robust than SVM once real-world sensor noise is present?",
    "What complementary failure modes do RF (tabular, noisy features) and LSTM (sequence, cold-start, data hunger) have, and how does an ensemble hedge them?",
    "How does the survey methodology (IEEE Xplore, PubMed, Springer, MDPI; 2016–2025) limit claims of completeness?",
    "Table II contrasts custom IoT hardware with consumer APIs. What deployment and cost argument follows from choosing APIs?",
    "Describe the academic-demo enrollment path: name, age, gender → next patient ID → 24-hour synthetic history → forecast.",
    "How is short-horizon forecasting (linear trend + HRI) different from the LSTM hypoglycaemia 30-minute lead time in Rigla et al. [15]?",
    "What does ‘simultaneous three-channel alert’ require of the software’s alert_system that a single toast notification does not?",
    "How are patient identifiers (P001–P500) and demographics used so the assistant cannot treat age and gender as ‘not specified’?",
    "Where in the pipeline would FHIR/HL7 mapping sit if this moved beyond an academic prototype?",
    # Solution (35–50) — cores 6,7,9 paraphrased
    "What limitation of Kajornkasirat [1] and Ekanath [2] alert designs does Gap 2 in Table III mark, and how does your demo close it on screen?",
    "Laranjo et al. [17] found engagement gains from conversational agents that never saw live wearables. How is your Ollama assistant architecturally different?",
    "What ethical boundary must you state before using the phrase ‘early disease detection’ in a viva?",
    "If TensorFlow or LSTM weights are missing, how does the risk engine still produce an HRI, and why is that acceptable in a demo?",
    "How would you demonstrate a critical test patient (for example extreme tachycardia and desaturation) to show alerts and the assistant together?",
    "Kolasa and Grzybowski [20] cite large per-patient savings in high-income settings. Why must you not quote those rupee-for-dollar as your project’s result?",
    "What does PulseVanguard AI actually implement today versus what Table II lists as a full lifestyle-recommendation engine?",
    "How do you answer: ‘Can this system replace a physician?’ without undermining the contribution?",
    "Perez et al. [10] validated consumer ECG at population scale but without a prediction pipeline. How does your stack add the missing layer while remaining honest about ECG scope?",
    "Why is synthetic / educational data the correct substrate for this academic build, and what would a clinical translation study still require?",
    "Explain heuristic-floor logic: why must extreme vitals not be diluted by a weak Random Forest probability?",
    "How does injecting history text, trend summary, and forecast into the chat prompt operationalize ‘live-data virtual assistance’?",
    "Name one scalability claim (cloud-native, consumer-grade cost) and one constraint (localhost Ollama, stubbed OAuth) you must pair in the same answer.",
    "If asked for originality, which combination of HRI + three-channel alerts + live assistant is hard to find in any single reviewed paper?",
    "How would you map each software module—normalizer, risk_engine, alert_system, ollama_agent, patient_memory—onto Fig. 1’s stages?",
    "Close the problem–process–solution arc in 60 seconds: fragmented data, ensemble HRI pipeline, multi-recipient alerts plus grounded chat.",
]

SET_B = [
    # Problem (1–16) — cores reworded
    "In what sense does the paper argue that healthcare still ‘operates reactively,’ and what temporal blind spot does that create between visits?",
    "Which architectural failures—not merely weaker accuracy numbers—does the literature treat as the primary unmet need?",
    "Why do the authors claim most NCD deaths are preventable, and how does that public-health claim license an engineering project?",
    "Fragmented wearable data, missing personal baselines, and single-recipient alerts: why must these three gaps be closed together rather than sequentially?",
    "How would you use Deng et al.’s multi-sensor fusion finding to argue against a heart-rate-only prototype?",
    "Gupta and Sharma [3] achieved remote monitoring without an AI prediction layer. Why is that design now considered an incomplete solution?",
    "Ramesh and Kumar [4] reached 90% cardiac-arrest accuracy on ECG alone. What interoperability and wearable-link limitations keep that paper from solving your problem statement?",
    "What does ‘care delivery remains largely episodic’ mean for a patient with evolving hypertension over six silent months?",
    "How do consumer devices’ ability to track HR, SpO₂, sleep, activity, and temperature create an opportunity that clinical IT has underused?",
    "Why is battery life listed by Deng et al. [6] as a persistent field problem, and how does an API-first design partially sidestep custom-hardware drain?",
    "Patel and Singh [8] use blockchain for tamper-evidence. Why did you not adopt that stack for the academic framework, given latency and overhead limitations in Table I?",
    "What is wrong with treating ‘user satisfaction 95%’ in a small prototype [2] as evidence of clinical effectiveness?",
    "How does the absence of personalized baselines cause alert fatigue in otherwise healthy users with atypical resting physiology?",
    "Why is single-lead consumer arrhythmia detection [10][11][13] still not a substitute for a multi-parameter risk index?",
    "State the problem as a systems-engineering mismatch: abundant sensors, scarce integrated intelligence.",
    "Which search terms and databases would you cite if asked to reproduce the survey protocol?",
    # Process (17–35) — cores 4,5,10
    "Why is a single composite HRI preferable to firing independent alarms on each vital sign?",
    "What complementary strengths do tree ensembles and recurrent networks offer on longitudinal wearable series?",
    "Map the software modules (normalizer, risk engine, alerts, chat, history/forecast) onto the paper’s architecture table without reading from a slide.",
    "Describe feature comparison Table II: how does ‘consumer wearables via API’ change cost and deployment versus Arduino/MAX30102 builds?",
    "How is the Health Risk Index bounded to 0–100, and what do the 25/50/75 cut-points correspond to in the alert policy?",
    "Explain unit conversion and schema mapping as the concrete mechanism of interoperability in the demo.",
    "What training artefacts does train_model.py write, and which paths does config.py expect at runtime?",
    "How does patient_memory construct a six-hour forecast, and what physiological clamps prevent absurd projections?",
    "Why keep a heuristic HRI even after loading Random Forest weights?",
    "Walk through FastAPI endpoints a viva examiner might ask you to name: /patients, /dashboard, /chat, /predict.",
    "How are three alert channels (patient, caregiver, clinician) encoded as structured messages rather than free text?",
    "What does ‘keep_alive’ and a smaller local LLM imply about latency versus answer quality in the assistant?",
    "How would you justify 24-hour hourly history per synthetic patient as enough for a demo trend yet not enough for LSTM pretraining?",
    "Where do PhysioNet-style sepsis labels appear in the data model, and why must they not be presented as a diagnosed condition?",
    "Explain CORS and localhost serving of the website: why must the demo not be opened as a raw HTML file?",
    "How does enrollment allocate the next P-series identifier from the existing dataset rather than a disconnected U001 namespace?",
    "What privacy posture does the academic build take (local SQLite, gitignored secrets, synthetic identities) versus GDPR-grade PHI?",
    "How would a future FHIR Observation resource map onto your unified metric JSON?",
    "Describe the literature-to-pipeline link: each of the seven stages closes a gap named in Section III.",
    # Solution (36–50) — cores 6,7,9
    "Why is simultaneous patient–caregiver–clinician notification a research contribution rather than a notification-preference setting?",
    "Why is grounding the assistant in history plus forecast more than a chatbot novelty?",
    "Why is synthetic or educational data acceptable for this prototype but insufficient for a clinical-performance claim?",
    "How would you live-demo ‘moderate versus critical’ HRI so the panel sees different alert bundles?",
    "Torous et al. [19] reported ~26% adherence improvement from smartphone interventions. How do you use that as motivation without claiming your chat replicated the trial?",
    "What must the assistant never invent (age, gender, diagnosis), and how did you constrain the prompt to prevent template answers such as ‘not specified’?",
    "If Ollama is offline, what fallback does the stack provide, and why is that better than a spinning loader?",
    "How does the dashboard’s trend sentence (‘rising heart rate; declining SpO₂’) communicate process to a clinician-examiner?",
    "Ekanath et al. [2] advertised 3–5 second alerts. What latency budget does a local CPU LLM chat have, and how do you set expectations?",
    "Name the four critical test identities used to stress alerts (severe hypoxia/tachycardia versus bradycardia/hypotension) without claiming they are real patients.",
    "How would you extend three-channel alerts to WhatsApp or email without changing the HRI contract?",
    "What remains stubbed (OAuth to Fitbit/Google Fit) and what is fully executed (HRI, forecast, alerts, grounded chat)?",
    "How does ‘heuristic floor when HRI heuristic ≥ 75’ protect a viva demo from a silent RF that scores elevated vitals as low risk?",
    "Give a 45-second solution pitch: interoperable vitals → ensemble HRI → three-channel alerts → live assistant, academic not clinical.",
    "If asked ‘what would you publish next?’, which gap (prospective wearable trial, FHIR, true personal baseline) is the honest sequel?",
]

SET_C = [
    # Problem (1–16)
    "How does PulseVanguard convert a reactive care model into a continuous preventive loop on the demo screen?",
    "How does the implemented software explicitly address fragmentation, weak personalization, and single-recipient alerts?",
    "How would you defend the public-health rationale of this project to a humanities or MBA examiner in two minutes?",
    "A panelist says ‘wearables already exist.’ What problem remains after the watch is on the wrist?",
    "Why can clinically accurate sensors still fail public health if their data never leave vendor silos?",
    "How would you use Kaiser Permanente’s 12-hour early-warning result [16] as analogy—not proof—for wearable HRI?",
    "What does ‘no intelligent layer converting raw readings into predictions’ mean in terms of your missing-middle architecture?",
    "Why is hypertension/diabetes/CVD a better motivating cluster for this framework than a single rare disease?",
    "How does alert-to-one-stakeholder design create asymmetric information among patient, family, and clinician?",
    "What limitation of hospital-only EWS systems makes a home-wearable pathway necessary?",
    "Faust et al. [11] and Rajpurkar et al. [13] are ECG-centric. Why is your problem statement multi-parameter by design?",
    "If India-scale affordability is raised, how do consumer APIs and ‘zero custom hardware’ in Table III answer cost without pretending the demo is a national deployment?",
    "What is the difference between detecting atrial fibrillation and estimating a composite deterioration risk?",
    "How do you prevent the viva from sliding into ‘we built a chatbot’ instead of ‘we closed literature gaps’?",
    "State the failure of static virtual assistants [17] as a problem of stale context, not of NLP quality.",
    "Why must problem, process, and solution be told as one story rather than three disconnected slides?",
    # Process (17–34)
    "Using the live demo, narrate one patient’s path from search/enrollment to vitals, HRI, forecast, alerts, and a grounded question to the assistant.",
    "In the repository, what happens at runtime if LSTM weights are absent, and how do you still defend the ensemble story from the paper?",
    "How is HRI computed when both Random Forest output and a rule-based heuristic are available?",
    "Explain the 0–25 / 25–50 / 50–75 / 75–100 bands as a policy layer on top of a continuous score.",
    "How do you keep demographics authoritative in chat (CSV age/gender) when a small local model prefers ‘not specified’ templates?",
    "Describe linear-trend forecasting: slope on recent hours, clamps, then HRI on the projected vector.",
    "What is the unified ingest conceptually even when connector OAuth is a stub—why keep the schema?",
    "How does FastAPI serving website/demo.html plus JSON APIs illustrate the ‘dashboard and assistant’ stage of Table II?",
    "Walk through trigger_alerts: band → message list → in-memory history. What is simulated versus real delivery?",
    "Why mix RF probability with heuristic rather than taking argmax of either?",
    "How would you explain num_predict / num_ctx limits as engineering for viva-stable latency?",
    "Where do P002 / P010 / P050 / P100 sit as process tests of the alert policy?",
    "How is the literature’s ‘personalized baseline’ only partially realized in the current heuristic-plus-population bands?",
    "What logging/privacy rule applies to API keys in user_store, and why mention it if connectors are stubbed?",
    "How does Streamlit frontend/app.py relate to the marketing site—two views, one risk engine?",
    "Explain why opening file://demo.html breaks API fetches and why 127.0.0.1:8000 is required.",
    "How would you add a new vital (for example HRV) with the least disruption to HRI, alerts, and chat?",
    "Map Gap 1–7 in Table III to a yes/no of ‘implemented in code today’ versus ‘proposed in paper’.",
    # Solution (35–50)
    "Walk through what a critical HRI produces in the demo alert pane across the three channels.",
    "How does the Ollama agent ingest patient context, and which disclaimer must accompany every reply?",
    "How would you answer: ‘Is this a medical device?’ in one precise sentence?",
    "Show how a named patient (not only an ID) supports the engagement story Torous and Laranjo sketched.",
    "What would a clinician still need to see that the current HRI does not encode (meds, labs, comorbidities)?",
    "How do you present originality: first combination in this academic framing, not ‘first AI in healthcare’?",
    "If the RF model scores near zero on obviously extreme vitals, what software safeguard preserves a critical band for the demo?",
    "How would you evaluate the assistant qualitatively in viva (grounding, brevity, no diagnosis) without a BLEU score?",
    "What production steps separate this repo from the cloud-native free-tier vision in Gap 7?",
    "Give the solution as a closed loop: measure → normalize → score → notify three roles → explain with live context.",
    "How does the website copy (collect, enroll, normalize, predict, alert) stay consistent with Section V?",
    "Why must lifestyle/diet recommendations remain modest or absent in the live build even though the paper lists a lifestyle engine?",
    "How would a prospective study be designed if this left the academic sandbox (endpoints, consent, comparator)?",
    "What single screen sequence best proves you solved the three literature gaps, not just plotted a chart?",
    "End with limitations: synthetic data, local LLM, stubbed vendors, HRI not a diagnosis—then restate the architectural contribution.",
    "In 70 seconds, restate problem (reactive, siloed, single alert), process (seven stages, RF+heuristic HRI), solution (demo: forecast, three channels, grounded assistant).",
]


def _set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_question(doc: Document, n: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(f"{n}. ")
    _set_run_font(r, size=11, bold=True, color=RGBColor(0x3D, 0x4B, 0xC8))
    r2 = p.add_run(text)
    _set_run_font(r2, size=11)


def main() -> None:
    assert len(SET_A) == 50 and len(SET_B) == 50 and len(SET_C) == 50

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Presentation & Viva Question Bank")
    _set_run_font(r, name="Georgia", size=22, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run(
        "An Interoperable Wearable Health Data Framework for Real-Time Predictive\n"
        "Health Monitoring and Early Disease Risk Detection"
    )
    _set_run_font(r, name="Georgia", size=12, italic=True, color=RGBColor(0x3D, 0x4B, 0xC8))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Akash Deshmukh · Aarya Thakare · Tanish Khedekar  |  Guide: Dr. Nilima Mohite\n"
        "JSPM’s Rajarshi Shahu College of Engineering  ·  AY 2025–26\n"
        "150 questions in three interchangeable sets of 50  ·  Grounded in the IEEE survey and the PulseVanguard / wearable-health-framework software"
    )
    _set_run_font(r, size=10, color=RGBColor(0x4A, 0x50, 0x78))

    note = doc.add_paragraph()
    r = note.add_run("How to use these sets. ")
    _set_run_font(r, size=11, bold=True)
    r = note.add_run(
        "Each set of 50 is a complete viva: it can carry problem statement, process, and solution on its own. "
        "About ten core ideas (reactive care, three literature gaps, WHO/NCD rationale, HRI bands, RF+LSTM/heuristic ensemble, "
        "three-channel alerts, live-data assistant, interoperability, academic/not-clinical boundary, seven-stage pipeline) "
        "reappear in all three sets, but the wording, angle, and cited paper change so a second panel will not hear a carbon copy. "
        "Questions are written to be a little demanding, yet fully answerable from the survey paper plus the running demo "
        "(synthetic data, FastAPI site, HRI, forecast, alerts, Ollama chat). Always state that this is an educational prototype, not a medical device."
    )
    _set_run_font(r, size=11)

    themes = doc.add_paragraph()
    r = themes.add_run("Suggested pacing inside every set of 50. ")
    _set_run_font(r, size=11, bold=True)
    r = themes.add_run(
        "Set A leans on literature and problem framing; Set B on architecture and pipeline; Set C on the live software. "
        "If you are asked to ‘explain the project,’ answer in three beats: (1) reactive, siloed, single-recipient care; "
        "(2) normalize → ensemble HRI → forecast; (3) three-channel alerts + assistant grounded in that patient’s history."
    )
    _set_run_font(r, size=11)

    sets = [
        (
            "Set A — Fifty questions",
            "Emphasis: problem statement, survey evidence, and why the framework exists. Still includes process and solution so one panel is enough.",
            SET_A,
        ),
        (
            "Set B — Fifty questions",
            "Emphasis: process, architecture, data flow, and model choices. Problem and solution questions are re-angled, not copied.",
            SET_B,
        ),
        (
            "Set C — Fifty questions",
            "Emphasis: solution, PulseVanguard demo, and honest limitations. Problem and process return in software-facing language.",
            SET_C,
        ),
    ]

    for title, blurb, questions in sets:
        doc.add_page_break()
        h = doc.add_paragraph()
        r = h.add_run(title)
        _set_run_font(r, name="Georgia", size=16, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
        b = doc.add_paragraph()
        r = b.add_run(blurb)
        _set_run_font(r, size=10, italic=True, color=RGBColor(0x4A, 0x50, 0x78))
        for i, q in enumerate(questions, start=1):
            add_question(doc, i, q)

    doc.add_page_break()
    h = doc.add_paragraph()
    r = h.add_run("Internal cue card — ten recurring themes (do not read this page aloud as a script)")
    _set_run_font(r, name="Georgia", size=14, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
    cues = [
        "Reactive vs preventive: visits are snapshots; NCDs evolve continuously; WHO ~74% of deaths from NCDs, many preventable with earlier detection.",
        "Three gaps: fragmented vendor data; no personal baseline; alerts to only one role. Largest unmet need is architectural.",
        "Interoperability: Mandl & Kohane [18] open APIs; Dias & Cunha [14] software layer lagging hardware; unified schema in code.",
        "HRI: composite 0–100; bands low <25, moderate <50, high <75, critical ≥75; better than isolated vital alarms.",
        "Models: Zhou RF>SVM under noise; Miotto/Rigla LSTM for time series; software uses RF + heuristic (LSTM optional if TF present).",
        "Alerts: literature is single-recipient; contribution is simultaneous patient + caregiver + clinician messages in alert_system.",
        "Assistant: Laranjo chatbots lacked live wearables; Ollama prompt gets vitals, demographics, history, forecast; educational disclaimer.",
        "Pipeline: devices → normalize → AI/HRI → alerts → dashboard/assistant (seven stages in the paper; modules in the repo).",
        "Not clinical: synthetic/demo data; not a diagnosis or medical device; Escobar/Perez are motivation, not our trial results.",
        "Demo proof: pick a critical synthetic patient, show HRI band, three alerts, ask the assistant a grounded question.",
    ]
    for i, c in enumerate(cues, start=1):
        p = doc.add_paragraph()
        r = p.add_run(f"{i}. ")
        _set_run_font(r, size=11, bold=True)
        r2 = p.add_run(c)
        _set_run_font(r2, size=11)

    path = Path(r"d:\IEEE_Survey_Paper_Health_Framework_QA_150.docx")
    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    main()
