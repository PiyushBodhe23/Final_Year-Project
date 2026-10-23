"""Mixed viva bank: 30 easy + 20 harder per set of 50, with spoken answers."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Each item: (level, question, answer)  — 3E + 2H, ten times = 30 easy / 20 hard

SET_A = [
    ("easy", "What’s the actual name of the project, in one breath?",
     "An Interoperable Wearable Health Data Framework for real-time predictive monitoring. On the website we branded the demo PulseVanguard AI."),
    ("easy", "Who built it, and who guided you?",
     "Akash Deshmukh, Aarya Thakare, and Tanish Khedekar — final-year E&TC at JSPM’s RSCOE, Pune. Guide is Dr. Nilima Mohite, AY 2025–26."),
    ("hard", "If someone says ‘you just made another dashboard,’ what’s the real problem you’re attacking?",
     "Clinic visits are snapshots. Chronic disease moves in the weeks in between. Watches already collect signals, but the data sits in Apple/Fitbit/Samsung silos, there’s no shared risk score, and a warning usually reaches only one person. We’re trying to close those three gaps in one teaching system."),
    ("easy", "What does WHO’s 74% figure actually refer to?",
     "About three-quarters of deaths worldwide come from non-communicable diseases — heart disease, diabetes, hypertension, chronic lung disease — and a lot of that is preventable if we catch trouble earlier."),
    ("easy", "In plain English, what are NCDs?",
     "Long-term conditions you don’t catch like an infection — diabetes, high BP, heart disease."),
    ("hard", "The survey names three technology gaps. Can you list them without looking at the slide?",
     "Fragmented wearable data, no personal baseline, and alerts built for a single recipient. The paper’s point is those are architecture problems, not ‘we need a slightly better neural net.’"),
    ("easy", "What is HRI, simply?",
     "Health Risk Index — one 0 to 100 score that summarises how worrying the current vitals look in the demo."),
    ("easy", "What are the four bands in the software?",
     "Low 0–25, moderate 25–50, high 50–75, critical 75–100."),
    ("hard", "Why isn’t a high heart rate, by itself, enough to shout ‘early warning’?",
     "Heart rate jumps with stairs, coffee, or anxiety. Without SpO₂, BP, temperature, trend, and context, you get false alarms. Deng’s review also says multi-sensor fusion beats a single channel. HRI is our attempt to combine several vitals."),
    ("easy", "How many papers did you review, and over what years?",
     "Twenty peer-reviewed papers, 2016 to 2025."),
    ("easy", "Name the five themes of the survey.",
     "Wearable IoT monitoring, ML risk prediction, interoperability, smart alerts, and virtual health assistants."),
    ("hard", "Mandl and Kohane [18] — what’s the one-line takeaway, and how does your normalizer answer it?",
     "They said the biggest digital-health barrier is that systems don’t share data, and they pushed open APIs. We don’t implement a national API, but the demo’s unified schema — who, when, metric, value, unit, source — is that idea at student scale: stop each watch speaking a private language."),
    ("hard", "What does ‘interoperability’ mean in this project?",
     "Different apps’ numbers get rewritten into one common format so one model can score them."),
    ("hard", "Which models does the paper propose for HRI?",
     "Random Forest plus LSTM as an ensemble, and we keep simple medical-range rules as backup."),
    ("hard", "Zhou liked Random Forest; Miotto and Rigla liked LSTM. Why both, not SVM alone?",
     "Zhou saw RF holding up better when wearable noise is real. LSTM is for sequences — your last 24 hours, not one reading. SVM was fine in lab BP/HR papers but weaker once noise showed up. So the story is: tables + time, not one fashionable classifier."),
    ("easy", "What’s a three-channel alert?",
     "Same event explained to the patient, a caregiver, and a clinician together — not just one of them."),
    ("easy", "What did older papers’ alerts usually do?",
     "Ping either staff or family. Almost never all three at once. That’s Gap 2 in our table."),
    ("hard", "Why is ‘one recipient’ a safety issue, not a UX preference?",
     "If only the patient sees it, they may ignore it. If only the ICU sees it, the family is blind. If only WhatsApp-to-caregiver fires, the clinician has no trail. Simultaneous channels are how we stop that information gap — even if, in the demo, delivery is simulated on screen."),
    ("easy", "Is this a medical device?",
     "No. College demo, synthetic data, not a diagnosis, doesn’t replace a doctor."),
    ("easy", "Where do we run the demo?",
     "http://127.0.0.1:8000 for the site, /demo for the live board. Server has to be on."),
    ("hard", "Dias and Cunha talked about a huge software gap. Where do you sit relative to that?",
     "They said sensors got good; the missing product is the layer that aggregates and thinks. We didn’t fill a billion-dollar market. We built a small, honest version of that missing layer: normalize, score, alert, explain."),
    ("easy", "What vitals does the demo actually use?",
     "Heart rate, SpO₂, breathing rate, temperature, systolic/diastolic BP, and steps."),
    ("easy", "What’s in the 500-patient file?",
     "Synthetic people P001–P500, with names, age, gender, and roughly a day of hourly vitals."),
    ("hard", "Escobar / Kaiser cut ICU transfers with ML. Can you quote that as your result?",
     "No. That’s a hospital EHR early-warning system, 12-hour lead, about 18% fewer ICU transfers. We use it as motivation — ‘prediction in the workflow can matter’ — not as our wearable trial."),
    ("hard", "What’s the chatbot, practically?",
     "Ollama on the laptop. It reads this patient’s vitals, history, and a short forecast, then answers. Always with a disclaimer."),
    ("easy", "Why must the bot see live numbers?",
     "The chatbots Laranjo reviewed were basically static FAQs. Ours is supposed to talk about today’s HRI, not a frozen profile."),
    ("hard", "If TensorFlow isn’t installed, have you abandoned the paper’s ensemble?",
     "No. The engine still blends Random Forest — if weights exist — with the heuristic. LSTM is optional. For a viva that’s honest: the architecture allows all three; this laptop may run two."),
    ("easy", "How do you add a new person in the UI?",
     "Name, age, gender. They get the next ID and 24 hours of made-up history so charts work."),
    ("easy", "How do you search?",
     "Type a name or ID — Priya, P002 — and pick from the list. Clear the box if it says no matches."),
    ("hard", "Say the ethical line before you use the words ‘early disease detection.’",
     "This is not detecting disease in a real body. It’s a simulated risk index on synthetic vitals. ‘Early detection’ in the title is the research ambition of the survey, not a claim we clinically achieved."),
    ("easy", "What should every assistant reply include?",
     "That this is educational, not medical advice."),
    ("easy", "If they ask ‘can this replace doctors?’",
     "No. It’s a teaching monitor and explainer. Decisions stay with a clinician."),
    ("hard", "Walk the seven-stage pipeline like you’re explaining it to your parents.",
     "Watch collects. We tidy the numbers. AI plus rules make one score. That score is HRI. If it’s high we show three kinds of warning. A screen shows history and a small future sketch. A chat box explains it. That’s device to intervention, student version."),
    ("easy", "Name three wearables the paper talks about.",
     "Apple Watch, Samsung Galaxy Watch, Fitbit."),
    ("easy", "What is PulseVanguard AI?",
     "The face of the demo website sitting on our FastAPI app."),
    ("hard", "Kajornkasirat got high correlation with medical sensors. Why isn’t that the same as clinical decision support?",
     "Correlation means the sensor roughly matches a reference. Decision support means you change a care decision, with workflow, alerts, and evidence. They proved the gadget; they didn’t prove a care pathway. We shouldn’t pretend our HRI did either."),
    ("hard", "What does the forecast table show?",
     "About six hours ahead of HR, SpO₂, HRI, and band — for teaching, not prophecy."),
    ("easy", "What if Ollama is off?",
     "You still get a local fallback paragraph using the same vitals. We say the model is offline."),
    ("hard", "What’s actually original if 20 papers already exist?",
     "Not ‘first AI in healthcare.’ It’s putting HRI, three-channel alerts, and a live-data assistant in one academic stack that the literature kept splitting across different systems."),
    ("easy", "Which paper is the open-API one?",
     "Mandl and Kohane, 2016."),
    ("easy", "Synthetic data means what?",
     "Fake educational rows so we don’t use real hospital secrets in this project."),
    ("hard", "Kolasa talks big dollar savings. Do we claim that?",
     "No. High-income meta-analysis, not our Pune demo. Quote it as ‘the field thinks prevention can save money,’ never as our invoice."),
    ("easy", "Name one software limitation you’d volunteer.",
     "Fitbit/Google OAuth is stubbed. Chat can be slow on CPU. LSTM may be missing."),
    ("easy", "Why not open the HTML file directly?",
     "The browser then can’t call /patients or /chat. Use localhost port 8000."),
    ("hard", "If they want originality plus humility in 60 seconds, what’s the arc?",
     "Problem: late, siloed, one-person alerts. Process: unify, score HRI, forecast. Solution: demo with three alerts and a grounded bot. Limit: not a doctor, not a trial."),
    ("easy", "What databases did you search?",
     "IEEE Xplore, PubMed, Springer, MDPI, plus a couple of health-informatics sources."),
    ("easy", "What does ‘normalize’ mean on stage?",
     "Take messy vendor readings and make them look like one table."),
    ("hard", "Where would FHIR sit if this left college?",
     "After normalization — map our JSON into FHIR Observation resources so a hospital EHR could ingest it. Not in the current repo."),
    ("easy", "Torous in one number?",
     "About 26% better adherence in some smartphone health studies — motivation, not our trial."),
    ("easy", "20-second pitch?",
     "Watches already measure. We bring the numbers together, give one risk score, warn patient, family, and clinician, and let a small AI explain. Demo only."),
]

SET_B = [
    ("easy", "‘Reactive healthcare’ — say it in one line.",
     "We mostly wait until someone’s already unwell, instead of watching the trend quietly."),
    ("easy", "Why name Apple, Samsung, Fitbit at all?",
     "To show the sensors are already on wrists. The hole is software that thinks across brands."),
    ("hard", "The paper says care is ‘episodic.’ Paint that for a hypertension patient over six quiet months.",
     "They look fine at two OPDs. In between, BP and heart rate may be drifting. Nobody is holding a continuous strip. That’s the blind spot wearables could fill if the data weren’t trapped in an app."),
    ("easy", "What’s a personal baseline, simply?",
     "What’s normal for you. A marathon runner’s resting HR isn’t an 80-year-old’s. Fixed cut-offs punish both."),
    ("easy", "HRI bands again?",
     "Low, moderate, high, critical — 25, 50, 75 as the fences."),
    ("hard", "Why close fragmentation, baselines, and single alerts together, not one PhD at a time?",
     "A perfect model on one vendor’s CSV still doesn’t reach the family. A triple alert on raw noisy HR still false-alarms. The survey’s argument is the failure is the missing whole system."),
    ("easy", "Which files hold the 500 people?",
     "patient_history.csv, patient_demographics.csv, plus a training Excel snapshot."),
    ("easy", "IDs look like what?",
     "P001 to P500, then P501 if you add someone."),
    ("hard", "Gupta and Sharma monitored remotely with no AI layer. Why call that incomplete now?",
     "You can stream temperature and still have no idea if tomorrow looks worse. The missing middle is prediction plus a shared score — which is why we added HRI, not only charts."),
    ("easy", "What serves the website?",
     "FastAPI, uvicorn, port 8000, backend/main.py."),
    ("easy", "Where are the HTML pages?",
     "website/index.html and website/demo.html."),
    ("hard", "Why not blockchain like Patel and Singh, if the table mentions it?",
     "Their own limitation is latency and cost. For a college predictive demo that would slow us down without helping HRI or alerts. We kept local files and said so."),
    ("hard", "What does the risk engine mix?",
     "RF if we have weights, LSTM if TensorFlow is there, and always a rule-of-thumb heuristic, then a band."),
    ("easy", "LSTM missing — does the demo die?",
     "No. RF and/or rules still give HRI."),
    ("hard", "RF is good at noisy tables, LSTM at sequences. How does an ensemble hedge their failures?",
     "RF doesn’t really ‘remember’ yesterday’s slope. LSTM wants lots of history and can starve on 24 synthetic hours — cold start. Together, plus rules, we don’t bet the viva on one model having a bad day."),
    ("easy", "Are WhatsApp alerts real?",
     "No. We show three message cards. That’s enough to explain the design."),
    ("easy", "What does the /chat API need?",
     "Your question and which patient, so we can attach vitals and HRI."),
    ("hard", "Why a composite HRI instead of five separate beeps?",
     "Five beeps is alarm soup. One band drives one policy: who gets told, how strongly. Clinicians think in ‘how bad is this person,’ not five uncoordinated toasts."),
    ("easy", "What’s Ollama?",
     "Local LLMs on your PC. No paid cloud key for the demo chat."),
    ("easy", "Why is the first chat slow?",
     "The model is loading into RAM on CPU. Don’t smash Send."),
    ("hard", "Map normalizer, risk engine, alerts, chat, history onto the paper without reading Table II.",
     "Normalizer is the interoperability stage. Risk engine is AI plus HRI. Alerts are the three-channel stage. History/forecast is the longitudinal bit. Chat is the live-data assistant. Dashboard is the visualisation."),
    ("hard", "Unified JSON — what’s in it?",
     "User, time, metric type, value, unit, source."),
    ("easy", "Real Apple Health login?",
     "No. Stub. Scoring uses our CSV."),
    ("hard", "Consumer APIs versus Arduino MAX30102 kits — what’s the deployment argument?",
     "Custom kits mean you manufacture hardware. APIs mean people use watches they bought. Paper says low cost and scale. Our honesty: we didn’t finish real OAuth; we scored a dataset as if that layer existed."),
    ("easy", "Kajornkasirat in one line?",
     "IoT sensors, RF/MLP/SVM, good correlation, small sample."),
    ("easy", "Ekanath’s headline numbers?",
     "About 92% in a prototype, alerts in a few seconds, simulated data."),
    ("hard", "How is our six-hour linear forecast different from Rigla’s 30-minute hypo LSTM?",
     "Rigla had dense glucose streams and a diabetes-specific model. We fit a simple slope on recent hours and clamp to plausible vitals, then rescore HRI. It’s a teaching forecast, not a 30-minute clinical lead time."),
    ("easy", "Table III Gap 1 in spoken English?",
     "Watches don’t talk to each other. We answer with one schema."),
    ("easy", "Gap 2?",
     "Only one person gets the ping. We show three."),
    ("hard", "Keep heuristic even after RF loads — why?",
     "In testing, RF sometimes shrugged at crazy vitals. If the rule score is already ‘critical,’ we don’t let a sleepy forest drag HRI down. That’s the heuristic floor you can mention."),
    ("easy", "How do you start the server?",
     "Activate venv, cd into wearable-health-framework, uvicorn backend.main:app --reload --host 127.0.0.1 --port 8000."),
    ("easy", "What’s on the patient card?",
     "Name, ID, age, gender, latest HR, SpO₂, RR, temp, BP, steps."),
    ("hard", "Name endpoints an examiner might throw: patients, dashboard, chat, predict.",
     "/patients lists everyone. /patients/{id}/dashboard is the demo payload. /chat is the assistant. /predict is the forecast object. All on the same FastAPI app."),
    ("easy", "Critical test IDs?",
     "P002, P010, P050, P100 — we pushed their last hours to extreme simulated vitals."),
    ("easy", "SpO₂ means?",
     "Oxygen in the blood, as a percent. Low values push HRI up in the rules."),
    ("hard", "sepsis_label in the CSV — can you say the patient ‘has sepsis’?",
     "No. It’s a leftover experimental flag, not a hospital diagnosis. Same for synthetic names. Don’t dramatise."),
    ("easy", "Systolic vs diastolic?",
     "Top number when the heart squeezes, bottom when it rests. Both go into HRI."),
    ("easy", "Why steps?",
     "Context. Heart rate 110 while walking is not the same as 110 in bed. Also a model feature."),
    ("hard", "Privacy: are we GDPR-grade?",
     "No. Local SQLite, gitignored secrets, fake people. That’s student hygiene, not a certified PHI platform."),
    ("easy", "FastAPI is…?",
     "The Python web framework serving pages and JSON."),
    ("easy", "Disclaimer you always say?",
     "Educational demo, not medical advice."),
    ("hard", "If Ollama is down, why a fallback instead of a spinner forever?",
     "Viva continues. The fallback still recites HRI and vitals from the server. You look prepared, not broken."),
    ("hard", "Forecast vs HRI?",
     "HRI is now. Forecast is a short ‘if this trend continues’ sketch."),
    ("easy", "Backend language?",
     "Python."),
    ("hard", "Torous 26% — how do you use it without lying?",
     "‘Digital engagement can help adherence in published studies. Our bot is in that spirit, grounded in vitals. We did not rerun their trial.’"),
    ("easy", "What you must never claim?",
     "That we diagnosed, treated, or validated on real Indian hospital patients."),
    ("hard", "How alerts grow with band?",
     "Low: mostly the patient. Moderate: add caregiver. High and critical: all three, stronger tone."),
    ("hard", "Stubbed vs real in one breath?",
     "Real: HRI, forecast, three alert texts, grounded chat, 500 synthetic lives. Stubbed: actual Fitbit/Google tokens. Local: Ollama, not a national cloud."),
    ("easy", "15-second process?",
     "Collect, tidy, score, three alerts, charts, chatbot."),
    ("easy", "If the list says no matches?",
     "Clear search, hard-refresh /demo, confirm uvicorn is running."),
]

SET_C = [
    ("easy", "Demo URL?",
     "http://127.0.0.1:8000/demo"),
    ("easy", "Landing page?",
     "http://127.0.0.1:8000 — PulseVanguard."),
    ("hard", "Someone says wearables already exist. What’s still broken after the watch is on?",
     "The clinic still doesn’t see a unified, predicted risk, and the family isn’t in the loop. The watch solved sensing, not the care pathway."),
    ("easy", "Problem in one sentence?",
     "Scattered watch data, late visits, warnings that don’t reach patient, family, and doctor together."),
    ("easy", "Solution in one sentence?",
     "Unify, score HRI, alert three roles, explain with a bot that sees those numbers."),
    ("hard", "How does the live screen show a shift from reactive to ‘watching ahead’?",
     "You don’t wait for a complaint. You pick a person, see now, a six-hour sketch, and alerts if the score is ugly — still simulated, but that’s the loop."),
    ("easy", "Which panel proves three-channel alerts?",
     "Alerts — patient, caregiver, clinician for the same HRI."),
    ("easy", "Which panel proves prediction?",
     "Forecast — plus hours, HR, SpO₂, HRI, band."),
    ("hard", "How do you prove the assistant isn’t a generic ChatGPT?",
     "Ask for this patient’s HRI and age. It should match the card. If it says ‘not specified,’ that’s a miss — we inject demographics on purpose."),
    ("easy", "Good question to ask the bot in viva?",
     "‘What’s the risk band and the recent trend?’"),
    ("easy", "Bad thing for the bot to do?",
     "Prescribe a drug or pretend it’s a licensed doctor."),
    ("hard", "Narrate one patient click-path like a demo script.",
     "Search Priya or P002. Card fills. Glance HRI. Scroll history. Forecast. Alerts if critical. Ask the bot why the band is high. Say the disclaimer out loud."),
    ("easy", "New user gets how many hours of history?",
     "Twenty-four simulated hours."),
    ("easy", "Genders in the add form?",
     "Female and male, same as the cleaned CSV."),
    ("hard", "If LSTM weights are absent, how do you still defend the paper’s ensemble story?",
     "The paper proposes RF+LSTM. The code path is ‘use what’s there, never crash.’ Heuristic plus RF is still an ensemble of sorts. Future work is enabling LSTM on a machine with TensorFlow."),
    ("easy", "patient_memory.py in one line?",
     "History, trend sentence, short forecast."),
    ("easy", "alert_system.py?",
     "Band in, three role-messages out."),
    ("hard", "HRI when both RF and rules exist — how do you say it without equations?",
     "We mix the forest’s probability with a rule score based on vital ranges. If rules already scream critical, that floor wins so a weak forest can’t hide it."),
    ("easy", "ollama_agent.py?",
     "Builds the prompt and calls the local model."),
    ("easy", "risk_engine.py?",
     "Turns vitals into HRI and a band."),
    ("hard", "Paper promised a lifestyle engine. Demo doesn’t diet-coach. How do you say that without looking unfinished-by-accident?",
     "‘In the architecture table it’s proposed. This build prioritised score, alerts, and Q&A. Lifestyle is future work.’ That’s cleaner than hiding it."),
    ("easy", "Where is the RF pickle?",
     "backend/models/weights/rf_model.pkl"),
    ("easy", "Train command?",
     "python backend/models/train_model.py from the project folder."),
    ("hard", "Gap 1–7: what’s in code today vs only in the paper?",
     "In code: unified vitals, HRI, three alerts, live chat, consumer-watch story, synthetic scale. Partial: true personal baseline. Paper-only-ish: full lifestyle engine, real cloud free-tier, real vendor OAuth."),
    ("hard", "Why both Streamlit and the website?",
     "Streamlit for us tinkerers; website for the audience. Same brain."),
    ("easy", "Can examiners use this on a real patient today?",
     "No."),
    ("hard", "What would a clinician still need that HRI doesn’t encode?",
     "Meds, labs, comorbidities, symptoms, ECG strips. HRI is a vital-based teaching score, not a full chart."),
    ("hard", "Mandl wanted what?",
     "More open, patient-accessible health data via APIs."),
    ("easy", "Why automatic alerts in the design?",
     "Literature says if people must remember to open a tool, they often don’t."),
    ("hard", "Is this a medical device? One precise sentence.",
     "No — it is an educational software prototype on synthetic data and is not a registered medical device."),
    ("hard", "Multi-sensor fusion in kitchen English?",
     "Don’t let one noisy sensor run the whole show."),
    ("easy", "Two vitals that often lift HRI?",
     "Wild heart rate and SpO₂ heading toward 90% or below, in our rule table."),
    ("hard", "If RF scores ~0 on a crashing patient, what saved the demo?",
     "The heuristic floor. Rules still put you in critical so the alert pane isn’t empty during viva."),
    ("easy", "Ollama port? Website port?",
     "11434 and 8000."),
    ("easy", "How do you show 500 patients exist?",
     "Hard refresh the demo, open the dropdown, P001 with a name should be there."),
    ("hard", "What’s the fairest ‘what next?’ paper?",
     "A small prospective wearable study, or FHIR export, or a real personal baseline — not another chatbot skin."),
    ("easy", "Institute and department?",
     "JSPM’s Rajarshi Shahu College of Engineering, Pune. Electronics and Telecommunication."),
    ("easy", "How do you start a talk?",
     "Problem, idea, one live patient, limits."),
    ("hard", "Closed loop in one breath?",
     "Measure, tidy, score, tell three roles, explain with live context."),
    ("hard", "FHIR in one line?",
     "A standard for sharing clinical data. Next step, not built."),
    ("easy", "Why Python?",
     "Pandas, scikit-learn, FastAPI — fast for a student prototype."),
    ("hard", "End with limits, then contribution — can you?",
     "Limits: fake data, local bot, stubbed vendors, HRI isn’t a diagnosis. Contribution: one place that shows how the survey’s gaps could fit together."),
    ("easy", "Easy difference vs Kajornkasirat?",
     "They soldered sensors. We score a software dataset as if consumer APIs fed it."),
    ("easy", "Vs Ekanath?",
     "They leaned caregiver alerts. We show three roles and a live-data assistant."),
    ("hard", "70-second full story if the mic is dying?",
     "Care is late and siloed. Twenty papers, three gaps. We unify vitals, HRI with forest and rules, forecast a few hours, alert three people, chat with context. It’s a demo. Thank you."),
    ("easy", "Stuck chat — what do you do?",
     "Wait. Don’t queue five Sends. Refresh if needed."),
    ("easy", "display_name is?",
     "The human name next to P002 so we can search Priya, not only codes."),
    ("hard", "How do you stop the viva becoming ‘we made a chatbot’?",
     "Lead with gaps and HRI. Chat is the last stage, grounded in the score. If they only remember the bot, you under-sold the survey."),
    ("easy", "If the server’s down?",
     "Start uvicorn. The 500 people didn’t vanish."),
    ("easy", "Last line if nerves eat your vocabulary?",
     "Unify, score, alert three, explain — demo, not a doctor."),
]


def _font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_qa(doc: Document, n: int, level: str, question: str, answer: str) -> None:
    pq = doc.add_paragraph()
    pq.paragraph_format.space_after = Pt(2)
    pq.paragraph_format.space_before = Pt(8)
    r = pq.add_run(f"Q{n}. ")
    _font(r, size=11, bold=True, color=RGBColor(0x3D, 0x4B, 0xC8))
    r2 = pq.add_run(question)
    _font(r2, size=11, bold=True)
    tag = pq.add_run("  ·  " + ("easier" if level == "easy" else "a bit tougher"))
    _font(tag, size=9, italic=True, color=RGBColor(0x4A, 0x50, 0x78))

    pa = doc.add_paragraph()
    pa.paragraph_format.space_after = Pt(6)
    pa.paragraph_format.left_indent = Inches(0.15)
    r = pa.add_run("A. ")
    _font(r, size=11, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
    r2 = pa.add_run(answer)
    _font(r2, size=11)


def main() -> None:
    for s in (SET_A, SET_B, SET_C):
        assert len(s) == 50, len(s)
        assert sum(1 for x in s if x[0] == "easy") == 30
        assert sum(1 for x in s if x[0] == "hard") == 20

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.95)
    sec.right_margin = Inches(0.95)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Mixed viva bank — easy majority")
    _font(r, name="Georgia", size=22, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run(
        "150 questions with answers  ·  three sets of 50  ·  30 easier + 20 a bit tougher in each set\n"
        "Wording is conversational. Mix of the IEEE survey and the PulseVanguard demo."
    )
    _font(r, name="Georgia", size=12, italic=True, color=RGBColor(0x3D, 0x4B, 0xC8))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Akash Deshmukh · Aarya Thakare · Tanish Khedekar  |  Dr. Nilima Mohite\n"
        "JSPM’s RSCOE, Pune  ·  AY 2025–26"
    )
    _font(r, size=10, color=RGBColor(0x4A, 0x50, 0x78))

    n = doc.add_paragraph()
    r = n.add_run("How this is mixed. ")
    _font(r, size=11, bold=True)
    r = n.add_run(
        "Each set of 50 is 30 easier questions and 20 slightly harder ones, woven so it doesn’t feel like two separate papers. "
        "Tougher items still have short, speakable answers. Always add: educational prototype, synthetic data, not a medical device. "
        "The older ‘hard only’ and ‘easy only’ files on this disk are still there if you want a pure drill."
    )
    _font(r, size=11)

    blocks = [
        ("Set A — mixed (30 easy / 20 tougher)", "Story of the problem and the paper, with a little demo.", SET_A),
        ("Set B — mixed (30 easy / 20 tougher)", "How it works: models, files, alerts, APIs.", SET_B),
        ("Set C — mixed (30 easy / 20 tougher)", "Live demo, files, and what you say on stage.", SET_C),
    ]
    for title, blurb, items in blocks:
        doc.add_page_break()
        h = doc.add_paragraph()
        r = h.add_run(title)
        _font(r, name="Georgia", size=16, bold=True, color=RGBColor(0x1A, 0x1F, 0x4A))
        b = doc.add_paragraph()
        r = b.add_run(blurb)
        _font(r, size=10, italic=True, color=RGBColor(0x4A, 0x50, 0x78))
        for i, (level, q, a) in enumerate(items, start=1):
            add_qa(doc, i, level, q, a)

    path = Path(r"d:\IEEE_Survey_Paper_Health_Framework_QA_150_Mixed.docx")
    doc.save(path)
    print("wrote", path, "easy", sum(1 for s in (SET_A, SET_B, SET_C) for x in s if x[0] == "easy"),
          "hard", sum(1 for s in (SET_A, SET_B, SET_C) for x in s if x[0] == "hard"))


if __name__ == "__main__":
    main()
